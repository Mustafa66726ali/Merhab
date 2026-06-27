"""تصدير تقرير معلومات العضو (Excel / PDF)."""

from io import BytesIO

from django.utils import timezone

from apps.platforms.member_profile import (
    list_managed_events,
    list_member_messages,
    list_member_qr_scans,
)


def _member_header_rows(member: dict) -> list[tuple[str, str]]:
    return [
        ("الاسم", member.get("name", "")),
        ("البريد", member.get("email", "")),
        ("الدور", member.get("role_label", member.get("role", ""))),
        ("الحالة", member.get("status_label", "")),
        ("تاريخ الانضمام", member.get("joined_at", "")),
    ]


def export_member_xlsx(member: dict, profile: dict) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "معلومات العضو"

    ws.append(["تقرير معلومات العضو"])
    ws.append([])
    for label, value in _member_header_rows(member):
        ws.append([label, value])

    stats = profile["event_stats"]
    ws.append([])
    ws.append(["إحصائيات الفعاليات"])
    ws.append(["إجمالي الفعاليات", stats["total"]])
    ws.append(["فعاليات نشطة", stats["active"]])
    ws.append(["فعاليات مكتملة", stats["completed"]])

    sections = profile["sections"]
    user_id = member["id"]
    platform_id = member.get("platform_id")

    if sections.get("show_messages") and platform_id:
        ws.append([])
        ws.append(["الرسائل المرسلة", profile["messages_total"]])
        ws.append(["ID", "الضيف", "الفعالية", "التاريخ", "الوقت"])
        rows, _ = list_member_messages(user_id, platform_id)
        for row in rows:
            ws.append([row["id"], row["guest_name"], row["event_title"], row["date"], row["time"]])

    if sections.get("show_qr_scans") and platform_id:
        ws.append([])
        ws.append(["عمليات مسح QR", profile["qr_scans_total"]])
        ws.append(["ID", "الضيف", "الفعالية", "التاريخ", "الوقت"])
        rows, _ = list_member_qr_scans(user_id, platform_id)
        for row in rows:
            ws.append([row["id"], row["guest_name"], row["event_title"], row["date"], row["time"]])

    if sections.get("show_managed_events") and platform_id:
        ws.append([])
        ws.append(["الفعاليات المُدارة", profile["managed_events_total"]])
        ws.append([
            "ID", "الفعالية", "الحالة", "التاريخ", "إجمالي الضيوف",
            "مؤكد", "حضر", "اعتذر", "نسبة التأكيد %", "نسبة الغياب %",
        ])
        rows, _ = list_managed_events(user_id, platform_id)
        for row in rows:
            ws.append([
                row["id"],
                row["title"],
                row["status_label"],
                row["date"],
                row["guests_total"],
                row["confirmed"],
                row["attended"],
                row["declined"],
                row["confirmation_rate"],
                row["absence_rate"],
            ])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_member_pdf(member: dict, profile: dict) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError:
        raise RuntimeError("مكتبة reportlab غير مثبتة")

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=2 * cm, leftMargin=2 * cm)
    styles = getSampleStyleSheet()
    story = []

    story.append(Paragraph("Member Profile Report", styles["Title"]))
    story.append(Spacer(1, 12))

    header_data = [[label, value] for label, value in _member_header_rows(member)]
    header_table = Table(header_data, colWidths=[5 * cm, 10 * cm])
    header_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f3f0ff")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
    ]))
    story.append(header_table)
    story.append(Spacer(1, 16))

    stats = profile["event_stats"]
    stats_data = [
        ["Total Events", str(stats["total"])],
        ["Active Events", str(stats["active"])],
        ["Completed Events", str(stats["completed"])],
    ]
    stats_table = Table(stats_data, colWidths=[5 * cm, 5 * cm])
    stats_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#5b2eff")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    story.append(Paragraph("Event Statistics", styles["Heading2"]))
    story.append(stats_table)

    sections = profile["sections"]
    user_id = member["id"]
    platform_id = member.get("platform_id")

    def add_table(title: str, headers: list[str], rows: list[list]):
        story.append(Spacer(1, 12))
        story.append(Paragraph(title, styles["Heading2"]))
        if not rows:
            story.append(Paragraph("No data", styles["Normal"]))
            return
        data = [headers] + rows
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#5b2eff")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
        ]))
        story.append(table)

    if sections.get("show_messages") and platform_id:
        msg_rows, _ = list_member_messages(user_id, platform_id)
        add_table(
            f"Messages ({profile['messages_total']})",
            ["ID", "Guest", "Event", "Date", "Time"],
            [[r["id"], r["guest_name"], r["event_title"], r["date"], r["time"]] for r in msg_rows[:50]],
        )

    if sections.get("show_qr_scans") and platform_id:
        scan_rows, _ = list_member_qr_scans(user_id, platform_id)
        add_table(
            f"QR Scans ({profile['qr_scans_total']})",
            ["ID", "Guest", "Event", "Date", "Time"],
            [[r["id"], r["guest_name"], r["event_title"], r["date"], r["time"]] for r in scan_rows[:50]],
        )

    if sections.get("show_managed_events") and platform_id:
        ev_rows, _ = list_managed_events(user_id, platform_id)
        add_table(
            f"Managed Events ({profile['managed_events_total']})",
            ["ID", "Event", "Status", "Confirm %", "Absence %"],
            [
                [r["id"], r["title"], r["status_label"], r["confirmation_rate"], r["absence_rate"]]
                for r in ev_rows[:50]
            ],
        )

    story.append(Spacer(1, 12))
    story.append(
        Paragraph(
            f"Exported: {timezone.localtime().strftime('%Y-%m-%d %H:%M')}",
            styles["Normal"],
        )
    )

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()
